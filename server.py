import os
import logging
import uuid
import shutil
import subprocess
from pathlib import Path
from typing import List

from fastapi import FastAPI, File, UploadFile, BackgroundTasks, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from werkzeug.utils import secure_filename
from pptx import Presentation
from pdf2image import convert_from_path
from supabase import create_client, Client
from livekit import api

load_dotenv()

# --- CONFIGURATION ---
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger('ppt-backend')

app = FastAPI()

# 1. CORS Setup (CRITICAL FIX: Allow All Origins)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allows Codespaces, Vercel, and Localhost
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 2. Supabase Setup
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    logger.error("❌ Missing Supabase Keys in .env")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None

# 3. Directories
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# --- BACKGROUND TASK ---
def process_ppt_task(file_path: str, presentation_id: str):
    """
    Background task to convert PPT -> Images -> Upload to Supabase
    """
    temp_dir = os.path.join(UPLOAD_FOLDER, presentation_id)
    os.makedirs(temp_dir, exist_ok=True)
    
    try:
        filename = os.path.basename(file_path)
        logger.info(f"🚀 Starting processing for {filename} (ID: {presentation_id})")

        # A. Create Parent Record
        if supabase:
            supabase.table("presentations").insert({
                "id": presentation_id,
                "filename": filename
            }).execute()

        # B. Convert to PDF (LibreOffice)
        cmd = [
            "libreoffice", "--headless", "--invisible", "--nodefault", "--nofirststartwizard",
            "--convert-to", "pdf",
            "--outdir", temp_dir,
            file_path
        ]
        subprocess.run(cmd, check=True)
        
        # Find generated PDF
        pdf_files = [f for f in os.listdir(temp_dir) if f.endswith('.pdf')]
        if not pdf_files:
            raise Exception("PDF conversion failed. No PDF created.")
        
        pdf_path = os.path.join(temp_dir, pdf_files[0])

        # C. PDF -> Images
        images = convert_from_path(pdf_path)
        
        # D. Extract Text & Upload
        prs = Presentation(file_path)
        slides_data = []

        for i, (image, slide) in enumerate(zip(images, prs.slides), start=1):
            # Save locally first
            img_filename = f"slide_{i}.jpg"
            img_local_path = os.path.join(temp_dir, img_filename)
            image.save(img_local_path, "JPEG")

            # Upload to Supabase Storage
            public_url = ""
            if supabase:
                storage_path = f"{presentation_id}/{img_filename}"
                with open(img_local_path, 'rb') as f:
                    supabase.storage.from_("slides").upload(
                        path=storage_path,
                        file=f,
                        file_options={"content-type": "image/jpeg"}
                    )
                public_url = supabase.storage.from_("slides").get_public_url(storage_path)

            # Extract Text
            text_content = []
            if slide.shapes.title and slide.shapes.title.text:
                text_content.append(f"Title: {slide.shapes.title.text}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape != slide.shapes.title:
                    clean_text = shape.text.strip().replace('\n', ' ')
                    if clean_text:
                        text_content.append(clean_text)

            slides_data.append({
                "presentation_id": presentation_id,
                "slide_number": i,
                "image_url": public_url,
                "content": " ".join(text_content)
            })

        # E. Batch Insert
        if supabase and slides_data:
            supabase.table("slides").insert(slides_data).execute()

        logger.info(f"✅ Finished processing {len(slides_data)} slides.")

    except Exception as e:
        logger.error(f"❌ Error processing PPT: {e}")
    finally:
        # Cleanup
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir)
        if os.path.exists(file_path):
            os.remove(file_path)

# --- ROUTES ---

@app.get("/")
def health_check():
    return {"status": "healthy", "service": "FastAPI Backend"}

@app.post("/api/upload-ppt")
async def upload_ppt(
    background_tasks: BackgroundTasks, 
    file: UploadFile = File(...)
):
    if not file.filename.endswith('.pptx'):
        raise HTTPException(status_code=400, detail="Invalid file type. Only .pptx allowed.")

    presentation_id = str(uuid.uuid4())
    safe_filename = secure_filename(file.filename)
    save_path = os.path.join(UPLOAD_FOLDER, safe_filename)

    with open(save_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    background_tasks.add_task(process_ppt_task, save_path, presentation_id)

    return {
        "status": "processing",
        "presentation_id": presentation_id,
        "message": "Upload accepted. Processing started in background."
    }

@app.get("/api/connection-details")
def connection_details():
    livekit_url = os.getenv('LIVEKIT_URL')
    api_key = os.getenv('LIVEKIT_API_KEY')
    api_secret = os.getenv('LIVEKIT_API_SECRET')

    if not all([livekit_url, api_key, api_secret]):
        raise HTTPException(status_code=500, detail="Missing LiveKit keys")

    room_name = f"ppt_{os.urandom(4).hex()}"
    
    token = api.AccessToken(api_key, api_secret) \
        .with_identity(f"user_{os.urandom(4).hex()}") \
        .with_name("User") \
        .with_grants(api.VideoGrants(
            room_join=True,
            room=room_name,
            can_publish=True,
            can_subscribe=True
        )).to_jwt()

    return {
        "serverUrl": livekit_url,
        "roomName": room_name,
        "participantToken": token
    }
    }
